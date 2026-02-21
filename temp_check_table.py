import openpyxl
wb = openpyxl.load_workbook('plugins/excel-to-csv/skills/excel-to-csv/references/example.xlsx', data_only=True)
for sheet in wb.sheetnames:
    ws = wb[sheet]
    for tbl_name, tbl in ws.tables.items():
        print(f"Table: {tbl_name}, Ref: {tbl.ref}")
